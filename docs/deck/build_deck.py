"""Build docs/deck/GMAPS_Open_Agent_2026.pptx (python-pptx). Numbers are copied from README.md, docs/PAIRED_STATS.md and the
published model card in docs/model_card/; images are in docs/deck/img/ (LGA view screenshots, model-card charts). SIMULATED data."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

HERE = Path(__file__).resolve().parent
IMG = HERE / 'img'
BG, PANEL, INK, DIM, TEAL, AMBER, RED, BLUE = (RGBColor(0x0f, 0x11, 0x15), RGBColor(0x1d, 0x21, 0x29), RGBColor(0xee, 0xf1, 0xf5),
                                               RGBColor(0xa3, 0xad, 0xbd), RGBColor(0x2d, 0xd4, 0xbf), RGBColor(0xe2, 0xb4, 0x3a),
                                               RGBColor(0xff, 0x4d, 0x4f), RGBColor(0x60, 0xa5, 0xfa))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def slide(title, kicker=None, simulated=False):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    if kicker:
        text(s, kicker, 0.6, 0.35, 10, 0.4, 14, TEAL, bold=True)
    text(s, title, 0.6, 0.7, 12, 0.9, 34, INK, bold=True)
    if simulated:
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.6), Inches(0.35), Inches(2.2), Inches(0.38))
        b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x2a, 0x1d, 0x0b); b.line.color.rgb = AMBER
        tf = b.text_frame; tf.text = 'SIMULATED'; p = tf.paragraphs[0]; p.runs[0].font.size = Pt(13); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = AMBER
    return s


def text(s, t, x, y, w, h, size, color=INK, bold=False):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.text = t
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = 'Calibri'
    return tb


def bullets(s, items, x=0.6, y=1.75, w=12.1, h=5.4, size=20):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        head, _, rest = it.partition('|')
        r = p.add_run(); r.text = ('• ' + head) if not rest else ('• ' + head + ' '); r.font.size = Pt(size); r.font.bold = bool(rest); r.font.color.rgb = INK; r.font.name = 'Calibri'
        if rest:
            r2 = p.add_run(); r2.text = rest; r2.font.size = Pt(size); r2.font.color.rgb = DIM; r2.font.name = 'Calibri'
        p.space_after = Pt(10)
    return tb


def box(s, x, y, w, h, title, body, color):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = PANEL; b.line.color.rgb = color; b.line.width = Pt(2)
    tf = b.text_frame; tf.word_wrap = True; tf.text = title
    p = tf.paragraphs[0]; p.runs[0].font.size = Pt(17); p.runs[0].font.bold = True; p.runs[0].font.color.rgb = color
    p2 = tf.add_paragraph(); r = p2.add_run(); r.text = body; r.font.size = Pt(13); r.font.color.rgb = INK


def picture(s, name, x, y, w=None, h=None):
    s.shapes.add_picture(str(IMG / name), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


# 1 title
s = slide('GMAPS · Ground Movement Awareness and Prediction System')
text(s, 'Open Agent Hackathon 2026 · Wildcard [Tinkerer] · Franco Plathottathil', 0.6, 2.15, 12, 0.6, 22, DIM)
text(s, 'A controller-side cross-check for the tower: radio becomes typed clearances, a ledger tracks who holds each runway, '
        'and an agent loop warns when two recorded authorizations, or an open obligation, hold the same runway at the same time.',
     0.6, 3.1, 12, 1.6, 22, INK)
text(s, 'Decision support only. GMAPS never transmits; the controller decides. All traffic and positions shown are SIMULATED.', 0.6, 6.4, 12, 0.6, 16, AMBER)

# 2 problem
s = slide('The problem and who it is for', 'PROBLEM')
bullets(s, ['Runway incursions start with clearances.|A landing clearance and a crossing clearance on the same runway, both read back, both legitimate on their own.',
            'The tower holds it all in memory.|Many aircraft and vehicles on one frequency; a "stop" that nobody reads back is easy to lose.',
            'Target user: the tower (local) controller.|A second set of eyes that keeps the ledger of who is cleared for what, on each runway, and says why it is worried.',
            'What it is not:|not an autopilot, not a transmitter, not a claim about real-world outcomes. Everything here runs on simulated traffic.'])

# 3 solution
s = slide('The solution: a traced agent loop over a runway ledger', 'SOLUTION')
bullets(s, ['Transmission in →|grammar parser types it; what the grammar cannot type goes to a model, and the answer must pass a grounding rule.',
            'Ledger (shared state) →|authorized · acknowledged · reported · observed, per runway; a warning names both sides and the rule.',
            'Look-ahead →|on simulated positions it projects whether two movers meet at the crossing; when positions are stale it says "prediction unavailable", which is not "no conflict".',
            'Explainer →|Nemotron writes two plain sentences for the controller; the evaluator checks them before they are shown.',
            'Every step is one trace line|naming the role and the model that handled it.'])

# 4 architecture
s = slide('Agent workflow: roles that check each other', 'ARCHITECTURE')
xs = [0.5, 2.65, 4.8, 6.95, 9.1, 11.25]
for x, (t, b, c) in zip(xs, [('Planner', 'grammar first; residue → model; look-ahead only when a shared zone is involved', TEAL),
                             ('Executor', 'bounded tools only: parse, ingest, look-ahead, explain, post alert. No shell.', BLUE),
                             ('Evaluator', 'grounding rule on every model answer; one retry with the airport vocabulary; then escalate "needs readback"', AMBER),
                             ('Risk', 'SIM-1 look-ahead each second; stale feed → "prediction unavailable"', RGBColor(0xfb, 0x92, 0x3c)),
                             ('Explainer + Critic', 'Nemotron drafts two sentences; a second Nemotron call critiques them against the facts → one revision, else labelled ledger text', RGBColor(0xa7, 0x8b, 0xfa)),
                             ('Alert', 'display only, never a radio frequency', RED)]):
    box(s, x, 1.9, 1.95, 2.6, t, b, c)
text(s, 'Shared state: the runway ledger + runway memory (open obligations per runway) · trace: runs/traces/<run_id>.jsonl (role, provider, model_reported, latency, tokens, outcome, grounding)',
     0.5, 4.8, 12.3, 0.9, 16, DIM)
text(s, 'Real code paths, with tests: critique → revision loop, retry, escalation, fallback when the provider is down, replay of recorded traces in sample mode (no key, no network).', 0.5, 5.8, 12.3, 0.9, 16, INK)

# 5 two small models
s = slide('Two small models, one loop', 'MODELS')
picture(s, 'base_vs_adapter.png', 0.5, 1.7, w=6.3)
text(s, 'Fine-tuned 4B parser (pre-existing, published)', 7.1, 1.7, 5.8, 0.5, 19, BLUE, bold=True)
text(s, 'Qwen3-4B-Instruct-2507 + LoRA "gmaps-parser", local vLLM on the DGX Spark.\n'
        'Record exact 0.254 → 0.652 on 138 held-out Whisper-text transmissions; paired +39.9 points, 95 % CI [+31.9, +48.6], McNemar p = 5.6e-17.\n'
        'Blind scripted set (414): 0.365 → 0.906. Grammar alone: 0.377 / 0.773.\n'
        'huggingface.co/darkstar008/gmaps-parser-qwen3-4b-lora', 7.1, 2.2, 5.8, 2.6, 14, INK)
text(s, 'Nemotron 30B via Nightshift (new)', 7.1, 4.75, 5.8, 0.5, 19, TEAL, bold=True)
text(s, 'Residue parse and the two-sentence explanation of every warning, through Nightshift (OpenAI-compatible). '
        'model_reported is logged on every call. Live run on the Spark pending at the time this deck was written.', 7.1, 5.25, 5.8, 1.6, 14, INK)
text(s, 'SIMULATED data. Charts and numbers: docs/model_card/, docs/PAIRED_STATS.md.', 0.5, 6.9, 7, 0.4, 11, DIM)

# 6-8 product visuals
for name, title, cap in [('lga_SIM_E1_CROSS_DURING_ROLLOUT_40.png', 'LaGuardia view: crossing cleared while a plane is landing',
                          'Left: today (radio and eyes). Right: WITH GMAPS. R1 conflict on runway 04/22 at +28 s; look-ahead projects the overlap at the G crossing. '
                          'Real LGA layout (OpenStreetMap, ODbL), fictional traffic.'),
                         ('lga_SIM_E5_HOLD_NOT_SLOWING_44.png', 'Told to hold short, not slowing',
                          'The hold short was read back, so the ledger is quiet; the look-ahead raises it from motion: the vehicle cannot stop at the real G hold bar.'),
                         ('lga_SIM_E4_FEED_LOSS_46.png', 'Position feed drops out',
                          'The vehicle track goes stale: GMAPS withdraws the timed prediction and says "prediction unavailable"; the ledger warning stays open.')]:
    s = slide(title, 'PRODUCT', simulated=True)
    picture(s, name, 0.5, 1.75, w=12.3)
    text(s, cap, 0.5, 5.9, 12.3, 1.2, 16, DIM)

# 9 sponsor tech
s = slide('Sponsor technology: what is real today', 'SPONSOR TECH')
bullets(s, ['Meterless Nightshift:|NightshiftLLM (OpenAI-compatible, key from the environment, temperature 0, one retry on 5xx, typed "provider_unavailable"). '
            'It is the primary provider for every model call in the loop. Live Nemotron run on the Spark pending at time of writing.',
            'NVIDIA:|Nemotron 30B through Nightshift for the residue parse and the explainer; the fine-tuned 4B and the DGX Spark runtime are pre-existing.',
            'Meterless Engine (H-MEM as runway memory):|not integrated. A local runway memory with the same interface (put / open_items / close) runs and is tested; '
            'trace lines say engine: local.runway_memory.',
            'Zetaris:|not integrated. Table design (identifiers, transmissions, ledger) and an idempotent loader are ready; the agent uses a local store.',
            'Unique:|the cloud model never has the last word: a deterministic grounding rule accepts or rejects every answer, and the trace that proves which model '
            'answered is also the offline fixture judges replay.'], size=17)

# 10 stack
s = slide('Technical stack', 'STACK')
bullets(s, ['Agent loop:|Python, stdlib only (agent/open_agent_loop.py); traces as JSONL; tools/show_trace.py.',
            'Models:|Nemotron 30B via Meterless Nightshift; Qwen3-4B + LoRA on vLLM v0.27.1 (DGX Spark, GB10); Whisper large-v3-turbo.',
            'Domain core (pre-existing):|grammar parser, runway ledger with cited rules, SIM-1 look-ahead, airport identifier graphs.',
            'UI:|ui_scaffold/lga_scene.html: canvas, OSM geometry in the FAA diagram frame, AS-IS vs WITH GMAPS on one clock.',
            'Run it:|python3 demo/run_open_agent.py --provider sample (no key) · demo/record_open_agent.sh (live) · python3 tests/run_all.py'])

# 11 pre-existing vs new
s = slide('Pre-existing vs new', 'DECLARATION')
box(s, 0.5, 1.8, 6.0, 4.6, 'Pre-existing (commit e97a9a1 and earlier, 18 Oct)',
    'Grammar parser, ledger and rules, SIM-1 simulator and look-ahead, airport identifier data and diagrams, LocalLLM adapter, '
    'the fine-tuned 4B parser and its evaluation, the DGX Spark vLLM runtime, the original demo page.', AMBER)
box(s, 6.8, 1.8, 6.0, 4.6, 'New in the build window (e97a9a1..HEAD)',
    'NightshiftLLM provider and sample replay; TraceLogger; grounding rule for model answers; the six-role agent loop with retry, escalation and fallback; '
    'runway memory; data-store layer and Zetaris loader design; Nemotron-vs-4B comparison runner; the animated LaGuardia view; tests.', TEAL)

# 12 next
s = slide('Limits and next steps', 'FUTURE')
bullets(s, ['Limits:|simulated traffic and positions, synthetic phraseology; the look-ahead is computed on the fictional SIM-1 model; decision support only.',
            'Next:|Zetaris as the data layer for identifiers, transmissions and ledger records; Meterless H-MEM for runway memory across sessions.',
            'Next:|Nemotron-vs-4B comparison on the 138 held-out transmissions, same prompt and verifier.',
            'Next:|reviewed airport geometry, a real surveillance feed, and a qualified controller reviewing the scenarios.'])

out = HERE / 'GMAPS_Open_Agent_2026.pptx'
prs.save(out)
print('wrote', out, len(prs.slides), 'slides')
